
from pathlib import Path
import fitz
import pytesseract
from PIL import Image
from docx import Document
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "documents"

# Configure Tesseract OCR on Windows
pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)


def read_pdf(file_path):
    """Extract selectable text and OCR scanned PDF pages."""
    extracted_pages = []

    with fitz.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()

            # Use OCR if the page has little or no selectable text
            if len(text) < 30:
                pix = page.get_pixmap(
                    matrix=fitz.Matrix(2, 2),
                    alpha=False
                )
                image = Image.frombytes(
                    "RGB",
                    [pix.width, pix.height],
                    pix.samples
                )
                text = pytesseract.image_to_string(
                    image,
                    lang="eng"
                ).strip()

            if text:
                extracted_pages.append(
                    f"Page {page_number}:\n{text}"
                )

    return "\n\n".join(extracted_pages)


def read_docx(file_path):
    """Extract paragraphs and table contents from Word documents."""
    doc = Document(file_path)
    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)


def read_xlsx(file_path):
    """Extract cell contents from Excel workbooks."""
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )
    parts = []

    try:
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]
                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.close()

    return "\n".join(parts)


def read_image(file_path):
    """Extract text from an image using OCR."""
    with Image.open(file_path) as image:
        return pytesseract.image_to_string(
            image.convert("RGB"),
            lang="eng"
        ).strip()


def read_document(file_path):
    """Select a reader based on the file extension."""
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return read_pdf(file_path)
    elif extension == ".docx":
        return read_docx(file_path)
    elif extension == ".xlsx":
        return read_xlsx(file_path)
    elif extension in {
        ".png", ".jpg", ".jpeg", ".tiff", ".bmp"
    }:
        return read_image(file_path)
    else:
        raise ValueError(
            f"Unsupported file format: {extension}"
        )


def read_all_documents():
    """Read supported documents from the documents folder."""
    DOCS_DIR.mkdir(exist_ok=True)
    results = {}

    supported_extensions = {
        ".pdf", ".docx", ".xlsx",
        ".png", ".jpg", ".jpeg", ".tiff", ".bmp"
    }

    for file_path in sorted(DOCS_DIR.iterdir()):
        if not file_path.is_file():
            continue

        if file_path.name.startswith("~$"):
            continue

        if file_path.suffix.lower() not in supported_extensions:
            continue

        try:
            results[file_path.name] = read_document(file_path)
        except Exception as error:
            results[file_path.name] = f"Error: {error}"

    return results


# Keep compatibility with the earlier PDF-reading code
def read_all_pdfs():
    return {
        name: text
        for name, text in read_all_documents().items()
        if name.lower().endswith(".pdf")
    }


if __name__ == "__main__":
    documents = read_all_documents()

    if not documents:
        print("No supported documents found in the documents folder.")

    for filename, content in documents.items():
        print(f"\n{'=' * 50}")
        print(f"FILE: {filename}")
        print(f"{'=' * 50}")
        print(content[:5000])
